"""
The audit the owner signs: what the money adds up to after migration, stock
per warehouse against `Persediaan Barang`, and every flag the replay raised —
a workbook to check against B2Soft line by line, and a summary to read first.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

PLACEHOLDER = re.compile(r"^S?(ERSEN|YOUNGER|PRASS|REFUND|AJUST|ADJUST|ABCDEF)")
RP = "#,##0"


def _sheet(wb, title, header, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = RP
    for i, _ in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 18 if i > 1 else 34
    ws.freeze_panes = "A2"
    return ws


def write(xlsx: Path, md: Path, *, docs, L, stock, customers, suppliers, report_piutang: int, report_hutang: int) -> dict:
    rows = docs.rows
    flags = list(L.flags)
    for u in L.units.values():
        if not u.pseudo and PLACEHOLDER.match(u.serial):
            flags.append(("Serial bukan nomor seri Apple", u.serial, f"{u.product}, status {u.status}"))

    pending = sum(int(p["amount"]) for o in rows["orders"] for p in o.get("payments", []) if p["state"] == "PENDING")
    due = sum(int(p["amount"]) for o in rows["orders"] for p in o.get("payments", []) if p["state"] == "DUE")
    opening_c = sum(int(c.get("openingBalance", "0")) for c in rows["customers"])
    opening_s = sum(int(s.get("openingBalance", "0")) for s in rows["suppliers"])
    open_po = 0
    for po in rows["pos"]:
        billed = sum(int(i["unitCost"]) * i["qty"] for i in po["items"])
        open_po += max(0, billed - sum(int(p["amount"]) for p in po["payments"]))
    master_c = -sum(p.balance for p in customers)
    master_s = -sum(p.balance for p in suppliers)

    live = Counter()
    for u in L.units.values():
        if u.status in ("TERSEDIA", "QC"):
            live[(u.product, u.loc[3:])] += 1
    counted = Counter({k: q for k, (q, _) in stock.items() if q})
    status = Counter(u.status for u in L.units.values())
    topics = Counter(f[0] for f in flags)

    figures = [
        ("Faktur pembelian", len(rows["pos"])), ("Faktur penjualan", len(rows["orders"])),
        ("Retur", len(rows["returns"])), ("Pelanggan", len(rows["customers"])),
        ("Supplier", len(rows["suppliers"])), ("Salesman", len(rows["salesmen"])),
        ("Model", len(rows["models"])), ("Varian (SKU)", len(rows["skus"])),
        ("Unit (serial)", sum(1 for u in L.units.values() if not u.pseudo)),
        ("Unit aksesoris (serial semu)", sum(1 for u in L.units.values() if u.pseudo)),
        ("Riwayat unit (event)", sum(len(u.events) for u in L.units.values())),
        ("Stok TERSEDIA", status["TERSEDIA"]), ("Stok QC (gudang 04)", status["QC"]),
        ("Persediaan Barang B2Soft (unit)", sum(counted.values())), ("Stok setelah migrasi (unit)", sum(live.values())),
        ("Piutang marketplace (PENDING)", pending), ("Piutang pelanggan (DUE)", due),
        ("Saldo awal pelanggan", opening_c), ("Piutang total setelah migrasi", pending + due + opening_c),
        ("Master Pelanggan B2Soft (total saldo)", master_c), ("Laporan Piutang Dagang B2Soft", report_piutang),
        ("Hutang faktur terbuka", open_po), ("Saldo awal supplier", opening_s),
        ("Hutang total setelah migrasi", open_po + opening_s),
        ("Master Supplier B2Soft (total saldo)", master_s), ("Laporan Hutang Dagang B2Soft", report_hutang),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Ringkasan", ["Angka", "Nilai"], figures, money_cols=(2,))
    _sheet(wb, "Piutang", ["Pelanggan", "Saldo B2Soft", "Faktur terbuka", "PENDING (marketplace)", "DUE", "Saldo awal", "Total setelah migrasi"],
           [(b["name"], b["master"], b["invoices"], b["pending"], b["due"], b["opening"], b["invoices"] + b["opening"])
            for b in docs.balances if b["kind"] == "customer"], money_cols=range(2, 8))
    _sheet(wb, "Faktur belum lunas", ["Faktur", "Pelanggan", "Tanggal", "Total", "Dibayar", "Sisa", "Status", "Jatuh tempo"],
           [(o["code"], o["party"], o["at"], o["total"], o["paid"], o["open"], o["state"], o["due"]) for o in docs.open_sales],
           money_cols=(4, 5, 6))
    _sheet(wb, "Hutang", ["Supplier", "Saldo B2Soft", "Faktur terbuka", "Saldo awal", "Total setelah migrasi"],
           [(b["name"], b["master"], b["invoices"], b["opening"], b["invoices"] + b["opening"])
            for b in docs.balances if b["kind"] == "supplier"], money_cols=range(2, 6))
    keys = sorted(set(live) | set(counted))
    _sheet(wb, "Stok per gudang", ["Produk", "Gudang", "Persediaan B2Soft", "Setelah migrasi", "Selisih"],
           [(p, g, counted.get((p, g), 0), live.get((p, g), 0), live.get((p, g), 0) - counted.get((p, g), 0)) for p, g in keys])
    _sheet(wb, "Catatan", ["Topik", "Subjek", "Rincian"], sorted(flags))
    wb.save(xlsx)

    by_topic = defaultdict(list)
    for f in flags:
        by_topic[f[0]].append(f)
    lines = ["# B2Soft migration audit", "", f"Cutoff: {L.cutoff:%Y-%m-%d %H:%M} WIB.", "", "| Figure | Value |", "|---|---|"]
    lines += [f"| {k} | {v:,} |" for k, v in figures]
    lines += ["", "## Flags", "", "| Topic | Count | Example |", "|---|---|---|"]
    lines += [f"| {t} | {n} | {by_topic[t][0][1]} — {by_topic[t][0][2]} |" for t, n in topics.most_common()]
    md.write_text("\n".join(lines) + "\n")
    return dict(figures=dict(figures), topics=dict(topics))
